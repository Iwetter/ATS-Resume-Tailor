"""Milestone B: evidence-verified rewriting and formatting-preserving DOCX tailoring.

Uses a fictional master DOCX with deliberate formatting (custom margins,
fonts, a bold lead-in run, bullet style) and a realistic hardware internship."""
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
import pytest

from ats.docx_tailor import _signature, docx_to_pdf, export_tailored, key, map_items, pdf_check, section_layout, tailor_docx
from ats.evidence import verify_rewrite
from ats.extraction import extract_document
from ats.matching import extract_requirements, parse_resume, rank_projects, select_projects
from ats.models import Change
from ats.rewriting import apply_changes
from hunter.service import order_skills

HARDWARE_JOB = """Hardware Engineering Intern - Summer 2027
Responsibilities
Design and bring up printed circuit boards for sensor products.
Write embedded firmware in C++ for microcontrollers and debug I2C and SPI buses.
Verify digital logic for FPGA-based test equipment.
Required
Experience with schematic capture and PCB layout in KiCad or Altium.
Hands-on lab skills with an oscilloscope and multimeter.
Preferred
FPGA development in VHDL or Verilog."""
HEADINGS = {"EDUCATION", "EXPERIENCE", "PROJECTS", "SKILLS"}


def fictional_docx() -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(0.6)
    section.top_margin = section.bottom_margin = Inches(0.5)
    for line in open("tests/fixtures/fictional_resume.txt", encoding="utf-8").read().splitlines():
        if line in HEADINGS:
            run = doc.add_paragraph().add_run(line)
            run.bold, run.font.size, run.font.name = True, Pt(11), "Calibri"
        elif line.startswith("• "):
            run = doc.add_paragraph(style="List Bullet").add_run(line[2:])
            run.font.size, run.font.name = Pt(10), "Garamond"
        elif line.startswith("KiCad"):
            p = doc.add_paragraph()
            lead = p.add_run("Tools: ")
            lead.bold, lead.font.name = True, "Garamond"
            body = p.add_run(line)
            body.font.name, body.font.size = "Garamond", Pt(10)
        else:
            run = doc.add_paragraph().add_run(line)
            run.font.size, run.font.name = Pt(10), "Garamond"
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


@pytest.fixture
def docx_master():
    data = fictional_docx()
    return data, parse_resume(extract_document(data, "master.docx").text)


def bullet(master, text):
    return next(x for x in master.items if x.kind == "bullet" and text in x.text)


def tailored(master):
    """Real pipeline: rank projects for the job, keep the best two, order skills, apply a verified rewrite."""
    requirements = extract_requirements(HARDWARE_JOB)
    ranked = [g for g, _, _ in rank_projects(master, requirements)]
    chosen = ranked[:2][::-1]  # also exercise reordering
    draft = order_skills(select_projects(master, chosen), requirements)
    source = bullet(master, "four-layer")
    if source.group not in chosen:
        chosen = [source.group] + chosen[:1]
        draft = order_skills(select_projects(master, chosen), requirements)
    proposal = "Designed and built a four-layer KiCad PCB for a wearable sensor with 3.3 V power rails, then tested its I2C sensor communication."
    change = Change((source.id,), source.text, proposal, "Leads with PCB design and bring-up the job asks for.")
    return apply_changes(master, draft, (change,), {source.id}), change, chosen


# ---- evidence verification ------------------------------------------------------

def test_meaningful_rewrite_is_accepted_with_risk_flags(docx_master):
    _, master = docx_master
    source = bullet(master, "four-layer")
    result = verify_rewrite(Change((source.id,), source.text, tailored(master)[1].proposed_text, "x"), master)
    assert result.status == "review"
    assert any("Combines work from another line" in r for r in result.reasons)  # "tested" comes from the sibling bullet
    proposal = "Engineered a compact wearable sensor PCB in KiCad with 3.3 V power rails and built the prototype."
    flagged = verify_rewrite(Change((source.id,), source.text, proposal, "x"), master)
    assert flagged.status == "review" and any("New wording to confirm" in r and "compact" in r for r in flagged.reasons)


@pytest.mark.parametrize("proposal,why", [
    ("Designed, fabricated and assembled a four-layer wearable PCB in KiCad with 3.3 V power rails.", "stage of work"),
    ("Designed and built a four-layer wearable PCB in Altium with 3.3 V power rails.", "Technologies"),
    ("Designed and built a six-layer wearable PCB in KiCad with 3.3 V power rails at 2 GHz.", "Numbers"),
    ("Led a team that designed and built a four-layer wearable PCB in KiCad with 3.3 V power rails.", "stage of work"),
    ("Planned a four-layer wearable PCB in KiCad with 3.3 V power rails.", "finished"),
])
def test_unsupported_claims_are_rejected(docx_master, proposal, why):
    _, master = docx_master
    source = bullet(master, "four-layer")
    result = verify_rewrite(Change((source.id,), source.text, proposal, "x"), master)
    assert result.status == "rejected" and any(why in r for r in result.reasons), result


def test_assisted_work_keeps_its_qualifier():
    master = parse_resume("Alex Example\nPROJECTS\nRadar Board | 2026\n• Assisted with the layout of a radar PCB in Altium.")
    source = next(x for x in master.items if x.kind == "bullet")
    assert verify_rewrite(Change((source.id,), source.text, "Laid out a radar PCB in Altium.", "x"), master).status == "rejected"
    assert verify_rewrite(Change((source.id,), source.text, "Supported the radar PCB layout in Altium.", "x"), master).status == "review"


def test_approved_evidence_library_entry_can_support_a_tool():
    master = parse_resume("Alex Example\nPROJECTS\nRadar Board | 2026\n• Designed a radar PCB in Altium.")
    source = next(x for x in master.items if x.kind == "bullet")
    proposal = "Designed a radar PCB in Altium and simulated its filters in LTspice."
    assert verify_rewrite(Change((source.id,), source.text, proposal, "x"), master).status == "rejected"
    library = [{"text": "Simulated the radar filters in LTspice.", "applies_to": [source.id], "approved_at": "2026-10-23"}]
    assert verify_rewrite(Change((source.id,), source.text, proposal, "x"), master, library).status == "review"


# ---- formatting-preserving DOCX ---------------------------------------------------

def test_every_resume_line_maps_to_its_paragraph(docx_master):
    data, master = docx_master
    mapping = map_items(Document(BytesIO(data)), master)
    assert len(mapping) == len([x for x in master.items if key(x.text)])


def test_tailoring_edits_the_original_document_in_place(docx_master):
    data, master = docx_master
    final, change, chosen = tailored(master)
    out, report = tailor_docx(data, master, final)
    assert not report["problems"] and not report["format_changes"] and not report["unmapped"]
    assert section_layout(out) == section_layout(data)  # margins and page size untouched
    original, edited = Document(BytesIO(data)), Document(BytesIO(out))
    texts = [p.text for p in edited.paragraphs]
    assert change.proposed_text in texts and change.original_text not in texts
    # Omitted projects are gone, kept ones appear in the chosen order.
    dropped = {x.text for x in master.items if x.kind == "entry" and x.section == "Projects" and x.group not in chosen}
    assert dropped and not dropped & set(texts)
    entries = [t for t in texts if t in {x.text for x in master.items if x.kind == "entry" and x.section == "Projects"}]
    assert entries == [next(x.text for x in master.items if x.kind == "entry" and x.group == g) for g in chosen]
    # The edited bullet keeps style, font and size; the skills line keeps its bold lead-in run.
    para = next(p for p in edited.paragraphs if p.text == change.proposed_text)
    assert para.style.name == "List Bullet" and para.runs[0].font.name == "Garamond" and para.runs[0].font.size == Pt(10)
    skills = next(p for p in edited.paragraphs if p.text.startswith("Tools: "))
    assert skills.runs[0].bold and skills.runs[0].text == "Tools: " and not skills.runs[1].bold
    assert skills.text != next(p.text for p in original.paragraphs if p.text.startswith("Tools: "))  # job-relevant skills first
    # Unchanged paragraphs are byte-for-byte the same formatting.
    unchanged = [p for p in edited.paragraphs if p.text in {q.text for q in original.paragraphs}]
    by_text = {p.text: p for p in original.paragraphs}
    assert all(_signature(p) == _signature(by_text[p.text]) for p in unchanged)


def test_meaningfully_tailored_not_just_reordered(docx_master):
    data, master = docx_master
    final, _, _ = tailored(master)
    changed = [x for x in final.items if x.id in master.by_id and master.by_id[x.id].text != x.text]
    assert any(x.kind == "bullet" for x in changed)  # real wording change, not only order
    assert len(final.items) < len(master.items)  # less relevant material removed for space


def test_unmatched_content_is_reported_not_silently_changed(docx_master):
    data, master = docx_master
    doc = Document(BytesIO(data))
    target = bullet(master, "four-layer")
    for p in doc.paragraphs:
        if p.text == target.text:
            p.runs[0].text = "Edited in Word after the master was saved."
    out = BytesIO()
    doc.save(out)
    final, _, _ = tailored(master)
    _, report = tailor_docx(out.getvalue(), master, final)
    assert any("Could not find" in p for p in report["problems"]) and target.id in report["unmapped"]


def test_word_or_libreoffice_render_is_one_page_and_extractable(docx_master):
    data, master = docx_master
    pdf, converter = docx_to_pdf(data)
    if pdf is None:
        pytest.skip("No DOCX-to-PDF converter on this machine: " + converter)
    final, _, _ = tailored(master)
    result = export_tailored(data, master, final)
    assert result["ok"], result["issues"]
    assert result["pages"] == 1 and pdf_check(result["pdf"], final)[1] == []
