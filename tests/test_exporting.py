from dataclasses import replace
from io import BytesIO

from docx import Document
import pymupdf
import pytest

from ats.exporting import ExportError, export_resume, verify_pdf
from ats.matching import make_resume
from ats.models import Evidence


def test_one_page_pdf_and_docx_round_trip(master):
    bundle = export_resume(master)
    assert bundle.pages == 1 and 10 <= bundle.font_size <= 11
    assert verify_pdf(bundle.pdf, master)[0] == 1
    assert "Designed and built" in bundle.text and "40 mA" in bundle.text
    doc = Document(BytesIO(bundle.docx))
    assert not doc.tables
    assert len(doc.inline_shapes) == 0
    assert all(x.font.size.pt >= 10 for p in doc.paragraphs for x in p.runs if x.text)
    with pymupdf.open(stream=bundle.pdf, filetype="pdf") as pdf:
        assert len(pdf[0].get_links()) >= 2
        assert pdf[0].get_text().count("Alex Example") == 1


def test_overflow_is_rejected_without_silent_omission(master):
    many = master.items + tuple(Evidence(f"long{i}", "Designed and tested technical hardware using an oscilloscope and measured voltage for an engineering verification project in a laboratory environment.", "bullet", "Experience", "long") for i in range(90))
    with pytest.raises(ExportError, match="pages"):
        export_resume(make_resume(many, master.source_hash))
    assert len(master.items) < len(many)


def test_pdf_verification_detects_missing_selected_content(master):
    bundle = export_resume(master)
    different = make_resume(master.items + (Evidence("missing", "An essential verified accomplishment absent from the PDF.", "bullet", "Experience"),), master.source_hash)
    with pytest.raises(ExportError, match="missing"):
        verify_pdf(bundle.pdf, different)
