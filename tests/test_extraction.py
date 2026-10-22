from io import BytesIO
from unittest.mock import Mock
import json

from docx import Document
import pymupdf
import pytest
import requests

from ats.extraction import InputError, extract_document, fetch_job, job_text_from_html
from ats.exporting import export_resume


def test_docx_extraction(docx_bytes):
    result = extract_document(docx_bytes, "resume.DOCX")
    assert "3.68" in result.text and "3.3 V" in result.text
    assert "Designed and built a four-layer" in result.text
    assert result.pages is None


def test_pdf_extraction_and_units(master):
    pdf = export_resume(master).pdf
    result = extract_document(pdf, "resume.pdf")
    assert result.pages == 1
    assert "40 mA" in result.text and "VHDL" in result.text


@pytest.mark.parametrize("data,name", [(b"", "r.pdf"), (b"broken", "r.pdf"), (b"not a zip", "r.docx"), (b"some text", "r.txt")])
def test_missing_corrupt_unsupported(data, name):
    with pytest.raises(InputError):
        extract_document(data, name)


def test_blank_pdf_is_unreadable():
    with pymupdf.open() as doc:
        doc.new_page()
        raw = doc.tobytes()
    with pytest.raises(InputError, match="selectable text"):
        extract_document(raw, "r.pdf")


def test_encrypted_pdf():
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((30, 40), "This is a resume with engineering experience and other details.")
        raw = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret")
    with pytest.raises(InputError, match="Password"):
        extract_document(raw, "r.pdf")


def test_table_and_header_extraction_in_order():
    doc = Document()
    doc.add_paragraph("Alex Example and a sufficiently long engineering resume introduction")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "First project"
    table.cell(0, 1).text = "FPGA VHDL verification"
    doc.add_paragraph("After the table")
    doc.sections[0].header.paragraphs[0].text = "alex@example.com"
    buf = BytesIO()
    doc.save(buf)
    out = extract_document(buf.getvalue(), "r.docx")
    assert out.text.index("First project") < out.text.index("After the table")
    assert "alex@example.com" in out.text
    assert any("Tables" in w for w in out.warnings)
    assert any("Header/footer" in w for w in out.warnings)


def test_structured_job_html():
    description = "<h2>Required</h2><p>PCB hardware testing and FPGA verification.</p>"
    html = '<script type="application/ld+json">' + json.dumps({"@type": "JobPosting", "title": "Hardware Intern", "description": description}) + "</script>"
    out = job_text_from_html(html)
    assert "Required" in out and "PCB" in out and "<h2>" not in out


def test_failed_url_requests_pasted_text(monkeypatch):
    monkeypatch.setattr("ats.extraction._public_url", lambda u: None)
    monkeypatch.setattr(requests, "get", Mock(side_effect=requests.Timeout))
    with pytest.raises(InputError, match="Paste"):
        fetch_job("https://jobs.example.com/posting")


@pytest.mark.parametrize("url", ["file:///C:/private.txt", "http://127.0.0.1/job", "http://localhost/job", "http://169.254.169.254/latest/", "https://u:p@example.com/job"])
def test_nonpublic_urls_are_rejected(url):
    with pytest.raises(InputError):
        fetch_job(url)


def test_no_generic_page_fabrication():
    with pytest.raises(InputError, match="Paste"):
        job_text_from_html("<html><body>Please sign in.</body></html>")
