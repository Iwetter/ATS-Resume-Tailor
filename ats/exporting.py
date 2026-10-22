"""Single-column exports with a strict PDF round-trip and page-count gate."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import os
import re
import unicodedata
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import pymupdf
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether

from .models import ExportBundle, Resume


class ExportError(ValueError):
    pass


def visible_text(text: str) -> str:
    return unicodedata.normalize("NFKC", text).replace("–", "-").replace("—", "-").replace("−", "-").replace("\u00a0", " ")


def _fonts() -> tuple[str, str, str]:
    candidates = [
        (Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts", "arial.ttf", "arialbd.ttf", "Arial"),
        (Path("/usr/share/fonts/truetype/dejavu"), "DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVu Sans"),
        (Path("/Library/Fonts"), "Arial.ttf", "Arial Bold.ttf", "Arial"),
    ]
    for root, regular, bold, docx_name in candidates:
        if (root / regular).is_file() and (root / bold).is_file():
            if "ResumeFont" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ResumeFont", str(root / regular)))
                pdfmetrics.registerFont(TTFont("ResumeFontBold", str(root / bold)))
            return "ResumeFont", "ResumeFontBold", docx_name
    return "Helvetica", "Helvetica-Bold", "Arial"


LINK = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}|(?:https?://|www\.|linkedin\.com/|github\.com/)[^\s|<>]+")


def _url(value: str) -> str:
    return "mailto:" + value if "@" in value else value if value.startswith(("https://", "http://")) else "https://" + value


def _markup(text: str) -> str:
    text = visible_text(text)
    pieces, position = [], 0
    for match in LINK.finditer(text):
        value = match.group().rstrip(".,;)")
        pieces += [escape(text[position:match.start()]), '<link href="' + escape(_url(value), {'"': '&quot;'}) + '">' + escape(value) + '</link>']
        position = match.start() + len(value)
    pieces.append(escape(text[position:]))
    return "".join(pieces)


def _pdf(resume: Resume, size: float, gap: float) -> bytes:
    regular, bold, _ = _fonts()
    for item in resume.items:
        font = pdfmetrics.getFont(bold if item.kind in {"section", "entry"} else regular)
        if hasattr(font, "face") and hasattr(font.face, "charToGlyph"):
            unsupported = {c for c in visible_text(item.text) if ord(c) > 32 and ord(c) not in font.face.charToGlyph}
            if unsupported:
                raise ExportError("Export font cannot render: " + " ".join(sorted(unsupported)) + ". Use a supported font or transliterate explicitly in the source.")
    styles = {
        "body": ParagraphStyle("body", fontName=regular, fontSize=size, leading=size + 2, spaceAfter=gap, alignment=TA_LEFT, splitLongWords=True),
        "header": ParagraphStyle("header", fontName=regular, fontSize=size, leading=size + 2, spaceAfter=gap),
        "name": ParagraphStyle("name", fontName=bold, fontSize=17, leading=20, spaceAfter=3),
        "section": ParagraphStyle("section", fontName=bold, fontSize=11, leading=13, spaceBefore=6, spaceAfter=3, keepWithNext=True),
        "entry": ParagraphStyle("entry", fontName=bold, fontSize=size, leading=size+2, spaceAfter=gap, keepWithNext=True),
        "bullet": ParagraphStyle("bullet", fontName=regular, fontSize=size, leading=size+2, leftIndent=10, firstLineIndent=0, bulletIndent=0, spaceAfter=gap),
    }
    flow = []
    for i, item in enumerate(resume.items):
        style = "name" if i == 0 and item.kind == "header" else item.kind if item.kind in styles else "body"
        p = Paragraph(_markup(item.text), styles[style], bulletText="•" if item.kind == "bullet" else None)
        flow.append(p)
    buf = BytesIO()
    SimpleDocTemplate(buf, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=30, bottomMargin=30, title="Tailored Resume", author="", pageCompression=1).build(flow)
    return buf.getvalue()


def _hyperlinks(paragraph, text: str):
    pos = 0
    for match in LINK.finditer(text):
        value = match.group().rstrip(".,;)")
        paragraph.add_run(text[pos:match.start()])
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(_url(value), RT.HYPERLINK, is_external=True))
        run, run_text = OxmlElement("w:r"), OxmlElement("w:t")
        run_text.text = value
        run.append(run_text)
        link.append(run)
        paragraph._p.append(link)
        pos = match.start() + len(value)
    paragraph.add_run(text[pos:])


def _docx(resume: Resume, size: float, gap: float) -> bytes:
    _, _, font_name = _fonts()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(30/72)
    sec.left_margin = sec.right_margin = Inches(.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = font_name, Pt(size)
    normal.paragraph_format.line_spacing = Pt(size + 2)
    normal.paragraph_format.space_after = Pt(gap)
    for i, item in enumerate(resume.items):
        style = "Title" if i == 0 and item.kind == "header" else "Heading 1" if item.kind == "section" else "Normal"
        p = doc.add_paragraph(style=style)
        pf = p.paragraph_format
        pf.widow_control = False
        if item.kind in {"section", "entry"}:
            pf.keep_with_next = True
        if item.kind == "section":
            pf.space_before, pf.space_after = Pt(6), Pt(3)
        if style == "Title":
            pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(3), Pt(20)
        if item.kind == "bullet":
            pf.left_indent, pf.first_line_indent = Inches(10/72), Inches(-10/72)
            p.add_run("• ")
        _hyperlinks(p, visible_text(item.text))
        for run in p.runs:
            run.font.name = font_name
            run.font.size = Pt(17 if style == "Title" else 11 if item.kind == "section" else size)
            run.font.bold = style == "Title" or item.kind in {"section", "entry"}
            from docx.shared import RGBColor
            run.font.color.rgb = RGBColor(0, 0, 0)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _canonical(text: str) -> str:
    # PDF engines may insert whitespace within a wrapped link or punctuation.
    return re.sub(r"\s+", "", visible_text(text)).casefold()


def verify_pdf(pdf: bytes, resume: Resume) -> tuple[int, str]:
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        if len(doc) != 1:
            raise ExportError(f"Export is {len(doc)} pages. Select fewer projects or shorten approved bullets; nothing was silently removed.")
        text = "\n".join(page.get_text("text", sort=True) for page in doc)
        actual = _canonical(text)
        missing = [x.id for x in resume.items if _canonical(x.text) not in actual]
        if missing:
            raise ExportError("PDF text verification failed for evidence IDs: " + ", ".join(missing))
        for page in doc:
            for word in page.get_text("words"):
                x0, y0, x1, y1 = word[:4]
                if x0 < 25 or x1 > page.rect.width-25 or y0 < 20 or y1 > page.rect.height-20:
                    raise ExportError("PDF text extends beyond the safe page margins.")
    return 1, text


def export_resume(resume: Resume) -> ExportBundle:
    if not resume.items or not any(x.kind == "header" for x in resume.items):
        raise ExportError("A resume needs its original name/contact header.")
    pdf, chosen = b"", None
    # Bounded layout attempts only; never shrink below 10 pt or drop content.
    for size, gap in ((10.5, 2), (10, 1)):
        pdf = _pdf(resume, size, gap)
        with pymupdf.open(stream=pdf, filetype="pdf") as doc:
            if len(doc) == 1:
                chosen = (size, gap)
                break
    pages, text = verify_pdf(pdf, resume)
    if chosen is None:
        raise ExportError("The selected content does not fit one page.")
    docx = _docx(resume, *chosen)
    reopened = Document(BytesIO(docx))
    docx_text = _canonical("\n".join(p.text for p in reopened.paragraphs))
    missing = [x.id for x in resume.items if _canonical(x.text) not in docx_text]
    if missing:
        raise ExportError("DOCX text verification failed: " + ", ".join(missing))
    return ExportBundle(pdf, docx, pages, text, chosen[0])
